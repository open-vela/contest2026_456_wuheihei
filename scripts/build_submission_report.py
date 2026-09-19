#!/usr/bin/env python3
"""Build the Audio Sentinel openvela contest report from the official template."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


SANS_FONT = "/System/Library/Fonts/Hiragino Sans GB.ttc"
SERIF_FONT = "/System/Library/Fonts/Supplemental/Songti.ttc"
NAVY = "203548"
TEAL = "167D86"
LIGHT = "EAF2F3"
MID = "C9D6DC"
GRAY = "5D6870"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, serif: bool = True) -> None:
    font_name = "Arial Unicode MS"
    run.font.name = font_name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clear_container(element) -> None:
    for child in list(element):
        element.remove(child)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])


def configure_document(doc: Document) -> None:
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.15)
    section.right_margin = Cm(2.15)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header
    header.is_linked_to_previous = False
    clear_container(header._element)
    p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("2026 首届 openvela AI 硬件开发者大赛  ·  Audio Sentinel")
    set_run_font(r, 8.5, color=GRAY, serif=False)

    footer = section.footer
    footer.is_linked_to_previous = False
    clear_container(footer._element)
    p = footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("—  ")
    set_run_font(r, 8.5, color=GRAY, serif=False)
    page_run = p.add_run()
    set_run_font(page_run, 8.5, color=GRAY, serif=False)
    add_field(page_run, "PAGE")
    r = p.add_run("  —")
    set_run_font(r, 8.5, color=GRAY, serif=False)

    def ensure_style(name: str, style_type=WD_STYLE_TYPE.PARAGRAPH):
        try:
            return doc.styles[name]
        except KeyError:
            return doc.styles.add_style(name, style_type)

    normal = ensure_style("Normal")
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in (
        ("Title", 25, NAVY),
        ("Heading 1", 16, NAVY),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 11.5, NAVY),
    ):
        style = ensure_style(style_name)
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    bullet = ensure_style("List Bullet")
    bullet.font.name = "Arial Unicode MS"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    bullet.font.size = Pt(10.3)


def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 25, True, NAVY, serif=False)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        set_run_font(r, 13, True, TEAL, serif=False)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc: Document, text: str, bold_lead: str | None = None,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, 10.5, True, NAVY)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, 10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, 10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        r = p.add_run("• " + item)
        set_run_font(r, 10.3)


def add_table(doc: Document, rows: list[list[str]], widths: list[float] | None = None,
              header: bool = True) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[col_index])
            set_cell_border(cell)
            if header and row_index == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1.5)
            r = p.add_run(str(value))
            set_run_font(r, 9.2, bold=(header and row_index == 0),
                         color=NAVY if header and row_index == 0 else None,
                         serif=not (header and row_index == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def make_architecture_figure(path: Path) -> None:
    width, height = 1600, 720
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(SANS_FONT, 52)
    box_font = ImageFont.truetype(SANS_FONT, 33)
    small_font = ImageFont.truetype(SERIF_FONT, 25)
    draw.text((65, 35), "Audio Sentinel 三层架构与证据边界", font=title_font, fill="#203548")

    boxes = [
        (55, 155, 500, 590, "Python 训练与离线评测",
         ["ESC-50 官方五折", "92 维特征 / TinyMLP", "Float 与 INT8 指标", "噪声鲁棒性评测"]),
        (575, 155, 1025, 590, "Linux C 主机模拟器",
         ["当前 C 特征前端", "INT8 前向推理", "三秒决策逻辑", "仅作链路冒烟验证"]),
        (1100, 155, 1545, 590, "OpenVela R528S3 端",
         ["arecord 三秒采集", "LCD / 触摸 / LVGL", "本地阈值与告警", "真机量化指标未测"]),
    ]
    colors = ["#EAF2F3", "#F3F5F6", "#E8F3EF"]
    for idx, (x1, y1, x2, y2, title, lines) in enumerate(boxes):
        draw.rounded_rectangle((x1, y1, x2, y2), 24, fill=colors[idx], outline="#167D86", width=4)
        draw.text((x1 + 28, y1 + 35), title, font=box_font, fill="#203548")
        y = y1 + 120
        for line in lines:
            draw.ellipse((x1 + 34, y + 10, x1 + 48, y + 24), fill="#167D86")
            draw.text((x1 + 66, y), line, font=small_font, fill="#34444F")
            y += 67

    for x in (515, 1040):
        draw.line((x, 370, x + 45, 370), fill="#167D86", width=7)
        draw.polygon([(x + 45, 355), (x + 45, 385), (x + 65, 370)], fill="#167D86")
    draw.text((58, 640), "指标来源：Python 源录音级五折评测；C 模拟器与真机集成不替代准确率评测。",
              font=small_font, fill="#5D6870")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, quality=95)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 9, color=GRAY)


def build(template: Path, output: Path, figure: Path, metrics_path: Path) -> None:
    metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
    clean = metrics["conditions"]["clean"]
    make_architecture_figure(figure)

    doc = Document(str(template))
    configure_document(doc)

    doc.add_paragraph().paragraph_format.space_after = Pt(48)
    add_title(doc, "Audio Sentinel", "基于 OpenVela 的离线音频事件检测系统")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("2026 首届 openvela AI 硬件开发者大赛 · 技术报告")
    set_run_font(r, 14, True, NAVY, serif=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(35)
    add_table(doc, [
        ["项目", "内容"],
        ["队伍名称", "呜嘿嘿"],
        ["团队成员", "李炳霖、吴安琪"],
        ["选题方向", "AI 硬件产品创新（端侧离线音频感知原型）"],
        ["官方仓库", "https://github.com/open-vela/contest2026_456_wuheihei"],
        ["提交基线", "v10 冻结源码、模型与评测证据"],
    ], widths=[3.6, 11.8])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("报告依据当前仓库证据编写；未测量的真机指标不作推断")
    set_run_font(r, 9.5, color=GRAY)
    doc.add_page_break()

    add_heading(doc, "一、作品信息表", 1)
    add_table(doc, [
        ["项目", "说明"],
        ["作品名称", "Audio Sentinel：基于 OpenVela 的离线音频事件检测系统"],
        ["队伍名称", "呜嘿嘿"],
        ["团队分工", "李炳霖：主要开发与实验执行；吴安琪：协作测试与提交材料复核。"],
        ["选题方向", "AI 硬件产品创新"],
        ["硬件平台", "润芯微 OpenVela 适配开发板，R528S3 Gemini S1 配置"],
        ["代码状态", "完整源码位于官方仓库默认分支 dev-ai-contest-2026"],
    ], widths=[3.5, 11.9])

    add_heading(doc, "二、摘要", 1)
    add_body(doc, "Audio Sentinel 面向家庭安全、看护与隐私敏感场景，在 OpenVela 设备本地完成三秒音频采集、92 维特征提取、INT8 TinyMLP 推理和 LVGL 交互，可识别背景、咳嗽、玻璃破碎、婴儿哭声和犬吠五类声音。模型采用 ESC-50 官方五折验证，Float 总体准确率 83.85%，INT8 总体准确率 84.10%。本项目提交训练源码、C 推理源码、模型和复现实验结果；真机已完成代码适配与镜像构建，但未测量真机准确率、延迟、内存、功耗和长期稳定性。")
    p = doc.add_paragraph()
    r = p.add_run("关键词：")
    set_run_font(r, 10.5, True, NAVY)
    r = p.add_run("OpenVela；离线音频事件检测；TinyMLP；MFCC；INT8；边缘智能")
    set_run_font(r, 10.5)

    add_heading(doc, "三、技术报告", 1)
    add_heading(doc, "3.1 引言与问题定义", 2)
    add_body(doc, "家庭环境中的咳嗽、玻璃破碎、婴儿哭声和犬吠具有明确的安全或看护意义。云端语音服务能够提供强大模型，但持续上传原始音频会带来网络依赖、隐私暴露和额外功耗。本项目选择在端侧离线完成检测，仅输出类别、置信度和本地告警状态。")
    add_body(doc, "项目目标不是构造通用声学大模型，而是在资源受限的 OpenVela 设备上形成可审计、可部署的五类事件检测链路。设计约束包括轻量模型、纯 C 前向推理、固定 16 kHz 单声道输入、对静音误报进行抑制，以及把离线准确率与真实板集成状态严格分开。")
    add_bullets(doc, [
        "隐私：默认不上传原始音频，推理和交互均在本地完成。",
        "轻量：TinyMLP 仅 6,277 个参数，INT8 权重存储估算约 7,029 字节。",
        "可复现：按源录音分组执行 ESC-50 官方五折，避免同源窗口跨折泄漏。",
        "可核验：模型、指标 JSON、C 头文件、模拟器与板级源码随仓库提交。",
    ])

    add_heading(doc, "3.2 系统总体设计", 2)
    add_body(doc, "系统被明确划分为三层。Python 层负责数据准备、增强、五折训练与量化评估；Linux C 主机模拟器层负责确认当前 C 特征前端、INT8 推理和三秒决策链路能够编译执行；OpenVela 层负责真实板音频采集、触摸显示、LVGL 界面和启动集成。三个层次共享冻结模型，但证据口径不同。")
    doc.add_picture(str(figure), width=Inches(6.35))
    add_caption(doc, "图 1  三层架构与证据边界")
    add_table(doc, [
        ["层次", "主要职责", "可以声明", "不能替代"],
        ["Python", "训练、五折、噪声与量化评测", "模型准确率与各类召回", "真机运行效果"],
        ["C 主机模拟器", "编译当前 C 推理与决策逻辑", "链路可运行及门限行为", "独立测试集准确率"],
        ["OpenVela 真机", "麦克风、LCD、触摸、LVGL、启动", "代码适配和镜像构建", "未测的延迟、RAM、功耗"],
    ], widths=[2.8, 5.0, 4.1, 3.5])

    add_heading(doc, "3.3 数据、特征与模型算法", 2)
    add_heading(doc, "3.3.1 数据与验证协议", 3)
    add_body(doc, "数据来自 ESC-50。四类目标事件分别选择 cough、glass_break、baby_cry、dog_bark，其余样本用于构造 background。评测单位为源录音：同一源录音切出的所有一秒窗口始终位于同一折。最终汇总包含 390 条源录音、1,950 个一秒窗口，其中 background 230 条，其余四类各 40 条。")
    add_body(doc, "训练设置为 120 epochs、隐藏层 64、随机种子 42、增强比例 0.5、类别权重指数 0.7。报告同时给出总体准确率、宏平均召回率和各类别召回率，避免 background 数量较多掩盖事件类别表现。")

    add_heading(doc, "3.3.2 92 维特征", 3)
    add_body(doc, "每个一秒窗口先计算 40 个适合嵌入式实现的时域统计量，再追加 13 维 MFCC 的均值与标准差，以及 MFCC Delta 的均值与标准差，共 92 维。Python 与 C 端均采用 25 ms 帧长、10 ms 帧移、26 个 Mel 滤波器和 13 个 MFCC 系数。")

    add_heading(doc, "3.3.3 TinyMLP 与量化", 3)
    add_body(doc, "最终网络为 92→64→5 的单隐藏层 TinyMLP，隐藏层使用 ReLU，输出层使用 Softmax，共 6,277 个参数。Float32 参数及归一化常量占 25,844 字节；采用逐张量 INT8 权重和浮点尺度后估算为 7,029 字节，缩减 72.80%。端侧前向过程中按尺度即时反量化，减少权重存储，同时保持实现简单。")
    add_table(doc, [
        ["项目", "数值"],
        ["结构", "92 → 64 → 5"],
        ["参数量", "6,277"],
        ["Float32 参数与归一化数据", "25,844 B"],
        ["INT8 权重存储估算", "7,029 B"],
        ["估算缩减", "72.80%"],
        ["模型 SHA256", "5382765b0aeb0c0aa9f96f7f78b2672\n8acfcd20efd35ac7b045637889b50df62"],
    ], widths=[5.2, 10.2])

    add_heading(doc, "3.3.4 鲁棒性策略", 3)
    add_body(doc, "训练阶段使用增强缓存和类别权重；端侧增加静音能量门限 2.0e-5 与事件置信度门限 0.80。三秒片段最多分为三个一秒窗口，优先选择达到门限的最高置信度非背景窗口，否则按平均概率输出。该策略用于抑制安静环境误报，但会使低能量事件被判为 background。")

    add_heading(doc, "3.4 系统实现", 2)
    add_heading(doc, "3.4.1 Python 训练端", 3)
    add_body(doc, "`training/` 提供 ESC-50 导入、音频预处理、MFCC92 特征缓存、鲁棒增强、TinyMLP 训练、五折评测、INT8 量化和 C 头文件导出。冻结模型位于 `model/`，评测结果位于 `results/python_cv/`。数据集本身因体积与授权边界不随仓库提交。")

    add_heading(doc, "3.4.2 Linux C 主机模拟器", 3)
    add_body(doc, "主机模拟器直接编译 `app/audiodetect/` 中的 WAV 读取、C 特征提取、INT8 模型和三秒决策逻辑，不使用旧 Goldfish 镜像。其目的为确认部署路径与门限行为，不作为模型准确率评测。15 个便利样本出现 14/15 标签一致；其中 dog_bark_01.wav 的 PCM 全零，被静音门限判为 background。排除该无有效事件音频后 14/14，仍只能称为冒烟测试。")

    add_heading(doc, "3.4.3 OpenVela 端功能", 3)
    add_table(doc, [
        ["功能", "状态", "证据或边界"],
        ["3 秒麦克风采集", "已实现代码", "arecord，16 kHz / 16-bit / mono，写入临时 WAV"],
        ["C 特征与 INT8 推理", "已实现代码", "最多三个 1 秒窗口，静音门限与 0.80 阈值"],
        ["LVGL 本地界面", "已实现代码", "START/STOP/AUTO、状态、类别、置信度、能量、计数"],
        ["开机启动", "已集成", "rcS 延时后启动 audiodetect ui"],
        ["最近事件", "部分实现", "内存维护 3 条摘要；当前 UI 未创建历史列表控件"],
        ["HTTP/MQTT", "命令行可选路径", "未接入实时 UI，未做端到端真机验证"],
        ["关键词识别", "命令行 WAV 路径", "未接入实时 UI，未做本次真机验证"],
    ], widths=[3.2, 3.1, 9.1])
    add_body(doc, "当前通过 `arecord` 分段写入 WAV，而非连续 PCM/I2S 环形缓冲，因此片段间可能存在空隙；未实现重叠滑窗、多标签、事件起止定位、持久数据库、OTA、在线学习或模型在线更新。")
    add_body(doc, "板级材料中存在需要保留的复现风险：`board/.../configs/nsh/defconfig` 选择 T070S140B，而冻结参考 `nuttx_savedefconfig` 与 `openvela_active.config` 选择 ILI9341，LVGL LCD 后端也有差异。本次遵守“部署冻结”要求，不继续改板或重建镜像，因此报告只声明代码集成与既有镜像构建，不声称当前 overlay 已重新复现冻结镜像。")

    add_heading(doc, "3.4.4 自建 AI Coding Skill", 3)
    add_body(doc, "仓库提供 `skills/openvela-audio-validation/SKILL.md`，触发词包括“复现 Audio Sentinel v10”“核对 OpenVela 音频模型”“验证音频事件部署证据”。Skill 固化了分层审计、模型哈希、指标解释、C 模拟器验证、板级边界和输出格式，并明确禁止把 14/14 冒烟结果写成准确率。")
    add_body(doc, "该 Skill 是 AI Coding 开发流程 Skill，用于提高复现与审计一致性；它不是部署在 `/data/agent/skills/` 的 OpenVela ai_agent 运行时 Skill。当前项目没有接入 ai_agent，不把这两类 Skill 混同。")

    add_heading(doc, "3.5 测试方案与结果", 2)
    add_heading(doc, "3.5.1 Clean 五折结果", 3)
    add_table(doc, [
        ["指标", "Float32", "INT8"],
        ["正确源录音数", "327 / 390", "328 / 390"],
        ["总体准确率", f"{clean['float']['overall_accuracy']*100:.2f}%", f"{clean['int8']['overall_accuracy']*100:.2f}%"],
        ["宏平均召回率", f"{clean['float']['macro_recall']*100:.2f}%", f"{clean['int8']['macro_recall']*100:.2f}%"],
    ], widths=[5.5, 4.95, 4.95])
    per_rows = [["类别", "样本数", "Float Recall", "INT8 Recall"]]
    for label in ["background", "cough", "glass_break", "baby_cry", "dog_bark"]:
        f = clean["float"]["per_class"][label]
        i = clean["int8"]["per_class"][label]
        per_rows.append([label, str(f["support"]), f"{f['recall']*100:.2f}%", f"{i['recall']*100:.2f}%"])
    add_table(doc, per_rows, widths=[4.8, 3.0, 3.8, 3.8])
    add_body(doc, "本次五折数据上 INT8 比 Float 总体准确率高 0.26 个百分点，这是量化扰动造成的边界变化，不能推广为“量化必然提高精度”。可以声明的是：本次证据未观察到量化精度损失。")

    add_heading(doc, "3.5.2 噪声鲁棒性", 3)
    robustness = [["条件", "Float Acc", "Float Macro", "INT8 Acc", "INT8 Macro"]]
    for key, label in [("clean", "Clean"), ("snr_20db", "20 dB"), ("snr_10db", "10 dB"), ("snr_5db", "5 dB")]:
        c = metrics["conditions"][key]
        robustness.append([
            label,
            f"{c['float']['overall_accuracy']*100:.2f}%",
            f"{c['float']['macro_recall']*100:.2f}%",
            f"{c['int8']['overall_accuracy']*100:.2f}%",
            f"{c['int8']['macro_recall']*100:.2f}%",
        ])
    add_table(doc, robustness, widths=[2.9, 3.1, 3.1, 3.1, 3.1])
    add_body(doc, "10 dB 与 5 dB 下事件类别召回显著下降，模型明显偏向 background。例如 Float 5 dB 下 cough、baby_cry、dog_bark 的召回分别为 7.50%、10.00%、12.50%。因此不能宣传“强噪声下仍保持高事件识别率”。")

    add_heading(doc, "3.5.3 证据可复现性", 3)
    add_bullets(doc, [
        "仓库自检脚本确认必要源码、模型和结果文件存在，冻结模型 SHA256 匹配。",
        "指标 JSON 可由混淆矩阵重新计算，Float/INT8 结果及各类别 recall 一致。",
        "NPZ 模型可按仓库导出算法逐字节重建 Float 与 INT8 C 权重头文件。",
        "原始 ESC-50 数据、训练缓存和便利 WAV 未随仓库提交，完整重训仍需按 README 准备数据。",
        "`results/evidence_checksums.json` 保留两项未随仓库提交的历史证据哈希；本报告不据此声称已在当前仓独立复核全部历史执行日志。",
    ])

    add_heading(doc, "3.5.4 真机未测项目", 3)
    add_body(doc, "真实板已经完成源码适配与既有镜像构建，但本次提交阶段没有重新连接开发板，也没有新的串口日志、运行录像或测量数据。因此不提供真机五类准确率、端到端延迟、单次推理耗时、峰值内存、功耗、连续运行时长、异常恢复或不同麦克风/距离/混响条件下的定量结论。演示视频由团队另行录制，不属于本报告产物。")

    add_heading(doc, "3.6 AI 原生开发与可审计日志", 2)
    add_body(doc, "项目开发中使用 Codex 协助需求拆解、代码审查、脚本编写、证据核验和报告整理。为避免将 AI 参与写成不可核验的宣传，本仓库提交一段在正式 OpenVela `.repo` 工作区内运行的 Codex CLI 只读审计会话，并使用赛事官方日志 schema 和验证器检查。")
    add_table(doc, [
        ["项目", "如实说明"],
        ["AI Coding 代码占比", "未做逐行归因，不报告无依据百分比"],
        ["使用工具", "Codex Desktop 与 Codex CLI；官方提交日志来自 Codex CLI"],
        ["MCP", "未使用赛事专用 MCP；开发中使用本地文件、Git、SSH 与浏览器能力"],
        ["自建 Skill", "skills/openvela-audio-validation/SKILL.md（AI Coding 工作流 Skill）"],
        ["AI 日志", "71 条官方 schema 事件，官方 validate-log.py 校验 ALL OK"],
        ["Token", "项目累计量未可靠统计；本次审计会话 CLI 结束摘要为 139,482 tokens"],
    ], widths=[4.2, 11.2])
    add_body(doc, "当前赛事采集器 1.3.0 只识别旧式 `.message` transcript，而 Codex CLI 0.154 使用 `response_item`。仓库提供兼容导出脚本做确定性字段映射，随后由官方 `snapshot_core.py` 生成编号、manifest、脱敏和最终 JSONL；生成后的比赛日志未被手工编辑。更早的桌面会话因不在官方支持流程内且可能包含历史凭证，没有混入提交。")

    add_heading(doc, "3.7 总结与展望", 2)
    add_body(doc, "Audio Sentinel 已形成从真实数据、严格五折、轻量 TinyMLP、INT8 权重到 OpenVela C/LVGL 集成的完整源码链路。现有证据支持 Float 83.85%、INT8 84.10% 的离线源录音级五折结果，也支持当前 C 部署路径和板级集成状态。")
    add_body(doc, "项目的主要不足是强噪声事件召回下降、分段采集存在空隙、板级配置快照不一致，以及缺少真机性能与长期稳定性测量。若后续解冻部署，优先级应为统一可复现 defconfig、改为连续环形缓冲与重叠窗口、补充独立真机数据集和延迟/RAM/功耗证据；本次比赛提交不把这些未来工作写成已完成。")

    add_heading(doc, "四、评分点对照与材料索引", 1)
    add_table(doc, [
        ["评分关注点", "仓库证据", "状态"],
        ["完整源码", "app/、training/、scripts/、board/", "已提交"],
        ["模型与指标", "model/、results/python_cv/", "已提交"],
        ["端侧部署", "app/audiodetect/、board/", "代码与既有镜像证据；新真机量化未做"],
        ["自建 Skill", "skills/openvela-audio-validation/SKILL.md", "已提交，AI Coding 工作流"],
        ["AI Coding 日志", "logs/bilibilidev/", "官方 schema 校验通过"],
        ["技术报告", "docs/submission/", "DOCX 与 PDF"],
        ["演示视频", "由团队单独准备", "不在本次生成范围"],
    ], widths=[3.8, 6.3, 5.3])

    add_heading(doc, "附录 A：关键复现命令", 1)
    add_body(doc, "以下命令用于拉取完整工程并检查提交内容。完整训练需要先按 `training/README.md` 准备 ESC-50 数据与缓存。")
    commands = [
        "repo init -u https://github.com/open-vela/contest2026_456_wuheihei -b dev-ai-contest-2026 -m contest2026_456_wuheihei.xml",
        "repo sync -c -j8",
        "cd contest2026_456_wuheihei && ./scripts/verify_submission.sh",
        "./scripts/build_c_simulator.sh",
        "python3 ../.claude/skills/contest-log-collector/tools/validate-log.py logs/",
    ]
    for command in commands:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(command)
        r.font.name = "Menlo"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        r.font.size = Pt(8.7)
        r.font.color.rgb = RGBColor.from_string("34444F")

    add_heading(doc, "附录 B：声明", 1)
    add_body(doc, "本报告只使用当前官方仓库 v10 源码、模型与结果，不混入上一赛事的项目名称、仓库地址、模型结构或指标。报告中“已实现”指当前源码存在对应逻辑；“已验证”只用于存在可核验证据的项目；未测项目均明确披露。")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--figure", type=Path, required=True)
    parser.add_argument("--metrics", type=Path, required=True)
    args = parser.parse_args()
    build(args.template, args.output, args.figure, args.metrics)


if __name__ == "__main__":
    main()
